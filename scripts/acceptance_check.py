"""Synthetic acceptance operations using each repository's own installed runtime."""

from pathlib import Path
import csv
import hashlib
import secrets
import io
import json
import os
import shutil
import sqlite3
import subprocess
import sys
import time
import traceback
import zipfile

ROOT = Path(__file__).resolve().parents[1]
os.chdir(ROOT)
sys.path.insert(0, str(ROOT))
from app.service import Application
from localdesk.safety import InputError
from tests.fixtures import make_image, make_pdf

import argparse

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument(
    "--report-dir",
    type=Path,
    required=True,
    help="New ignored evidence folder for this run. Existing folders are preserved.",
)
args = parser.parse_args()
BASE = args.report_dir.resolve()
BASE.mkdir(parents=True, exist_ok=False)
WORK = BASE / "synthetic space unicode-ÃƒÂ©"
WORK.mkdir()
RESULTS = []
APP = None
PASSWORD = secrets.token_urlsafe(32)


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def record(name, func):
    started = time.monotonic()
    try:
        detail = func()
        result = dict(check=name, status="PASS", detail=detail)
    except Exception:
        result = dict(check=name, status="FAIL", traceback=traceback.format_exc())
    result["seconds"] = time.monotonic() - started
    RESULTS.append(result)
    (BASE / "results.json").write_text(
        json.dumps(RESULTS, indent=2, default=str), encoding="utf-8"
    )
    print(name, result["status"], flush=True)


def finish(action, body, expected="done"):
    ident = APP.common_post(action, body)["job_id"]
    deadline = time.monotonic() + 180
    while time.monotonic() < deadline:
        job = APP.jobs.get(ident)
        if job["status"] not in ("running", "queued"):
            assert job["status"] == expected, job
            return job["result"] if expected == "done" else job
        time.sleep(0.03)
    raise TimeoutError(action)


def output(artifact):
    return APP.download(artifact["path"]).read_bytes()


def restart():
    global APP
    APP.close()
    APP = Application(
        ROOT,
        BASE / "persistent-data",
        passphrase=PASSWORD,
    )


def file(name, content):
    path = WORK / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content.encode("utf-8") if isinstance(content, str) else content)
    return path


def office_fixtures():
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches

    doc = Document()
    doc.add_paragraph("Office canary qa@example.test public number 42")
    doc.core_properties.author = "PRIVATE_AUTHOR_CANARY"
    doc.save(WORK / "office.docx")
    book = Workbook()
    book.active.append(["Office canary qa@example.test", 42])
    book.save(WORK / "office.xlsx")
    book.close()
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = (
        "Office canary qa@example.test public number 42"
    )
    deck.save(WORK / "office.pptx")
    return [WORK / ("office" + ext) for ext in (".docx", ".xlsx", ".pptx")]


def recovery():
    def images():
        from PIL import Image
        from playwright.sync_api import sync_playwright
        import base64

        results = []
        with sync_playwright() as pw:
            browser = pw.chromium.launch()
            page = browser.new_page()
            try:
                for suffix in (".png", ".jpg", ".gif"):
                    path = WORK / ("red-pixels" + suffix)
                    image = Image.new("RGB", (48, 32), (240, 20, 30))
                    image.save(path)
                    image.close()
                    before = sha(path)
                    batch = finish("inspect", {"paths": [str(path)]})
                    result = finish("repair", {"id": batch["inspections"][0]["id"]})
                    raw = output(result["output"])
                    pixels = page.evaluate(
                        """async data => {
                        const image = new Image();
                        image.src = 'data:image/png;base64,' + data;
                        await image.decode();
                        const canvas = document.createElement('canvas');
                        canvas.width = image.width; canvas.height = image.height;
                        const context = canvas.getContext('2d');
                        context.drawImage(image, 0, 0);
                        return {width:image.width,height:image.height,
                                pixel:Array.from(context.getImageData(0,0,1,1).data)};
                    }""",
                        base64.b64encode(raw).decode(),
                    )
                    assert (pixels["width"], pixels["height"]) == (48, 32), pixels
                    assert all(
                        abs(a - b) <= 3
                        for a, b in zip(pixels["pixel"], [240, 20, 30, 255])
                    ), pixels
                    assert sha(path) == before
                    results.append({"recovery": result, "chromium_pixels": pixels})
            finally:
                browser.close()
        return results

    record(
        "RL-05 PNG JPEG GIF recovered pixels decoded independently by Chromium", images
    )

    def documents():
        paths = office_fixtures() + [
            make_pdf(WORK / "readable.pdf", text="Recovery public canary 42")
        ]
        results = []
        for path in paths:
            before = sha(path)
            batch = finish("inspect", {"paths": [str(path)]})
            result = finish("repair", {"id": batch["inspections"][0]["id"]})
            raw = output(result["output"])
            if path.suffix == ".pdf":
                from pypdf import PdfReader

                reader = PdfReader(io.BytesIO(raw))
                assert len(reader.pages) == 1
                assert "Recovery public canary 42" in reader.pages[0].extract_text()
            else:
                with zipfile.ZipFile(io.BytesIO(raw)) as z:
                    assert z.testzip() is None
                    assert b"42" in b"\n".join(z.read(n) for n in z.namelist())
            assert sha(path) == before
            results.append(result)
        return results

    record(
        "RL-04 recovered PDF DOCX XLSX PPTX readable contents and source hashes",
        documents,
    )

    def samples():
        sources = APP.get("examples", {})["files"]
        hashes = {p: sha(Path(p)) for p in sources}
        batch = finish("inspect", {"paths": sources})
        records = []
        for row in batch["inspections"]:
            report = row["report"]
            APP.get("hex", {"id": row["id"], "offset": 0})
            exported = APP.common_post("export", {"id": row["id"]})
            assert json.loads(output(exported)) == report
            if report["repair_available"]:
                result = finish("repair", {"id": row["id"]})
                assert result["original_unchanged"]
                assert (
                    hashlib.sha256(output(result["output"])).hexdigest()
                    == result["output_sha256"]
                )
                assert (
                    json.loads(output(result["report"]))["output_sha256"]
                    == result["output_sha256"]
                )
                records.append(result)
        assert hashes == {p: sha(Path(p)) for p in sources}
        return {"batch": batch, "repairs": records, "source_hashes": hashes}

    record("RL-01 included batch all terminal results report bytes and hashes", samples)

    def database():
        source = WORK / "binary.sqlite"
        db = sqlite3.connect(source)
        db.execute(
            "CREATE TABLE records(id INTEGER PRIMARY KEY, label TEXT, value BLOB)"
        )
        db.execute("INSERT INTO records VALUES(1,?,?)", ("known", b"\x00\xff\x01"))
        db.commit()
        db.close()
        before = sha(source)
        batch = finish("inspect", {"paths": [str(source)]})
        result = finish("repair", {"id": batch["inspections"][0]["id"]})
        db = sqlite3.connect(APP.download(result["output"]["path"]))
        assert db.execute("SELECT * FROM records").fetchall() == [
            (1, "known", b"\x00\xff\x01")
        ]
        assert db.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        db.close()
        assert sha(source) == before
        renamed = source.with_suffix(".renamed")
        source.rename(renamed)
        renamed.rename(source)
        return result

    record(
        "RL-06 SQLite independent table BLOB integrity and Windows handle cleanup",
        database,
    )

    def video():
        p = WORK / "synthetic.mp4"
        subprocess.run(
            [
                "ffmpeg",
                "-v",
                "error",
                "-f",
                "lavfi",
                "-i",
                "testsrc=size=160x120:rate=10",
                "-t",
                "1",
                "-pix_fmt",
                "yuv420p",
                str(p),
            ],
            check=True,
        )
        before = sha(p)
        batch = finish("inspect", {"paths": [str(p)]})
        result = finish("repair", {"id": batch["inspections"][0]["id"]})
        target = APP.download(result["output"]["path"])
        probe = json.loads(
            subprocess.check_output(
                [
                    "ffprobe",
                    "-v",
                    "error",
                    "-show_streams",
                    "-show_format",
                    "-of",
                    "json",
                    str(target),
                ]
            )
        )
        assert probe["streams"][0]["codec_type"] == "video"
        assert float(probe["format"]["duration"]) >= 0.9
        subprocess.run(
            ["ffmpeg", "-v", "error", "-i", str(target), "-f", "null", "-"], check=True
        )
        assert sha(p) == before
        return {"recovery": result, "independent_decode": probe}

    record(
        "RL-07 synthetic remux streams duration full independent frame decode", video
    )


def shared():
    def persist():
        APP.store.set("qa_canary", "PERSISTED_SYNTHETIC_CANARY")
        backup = APP.common_post("backup", {})
        raw = output(backup)
        assert b"PERSISTED_SYNTHETIC_CANARY" not in raw
        restart()
        assert APP.store.get("qa_canary") == "PERSISTED_SYNTHETIC_CANARY"
        from localdesk.storage import Store

        restored = BASE / "restored.vault"
        restored.write_bytes(raw)
        store = Store(restored, PASSWORD)
        assert store.get("qa_canary") == "PERSISTED_SYNTHETIC_CANARY"
        store.close()
        before = sha(restored)
        try:
            Store(restored, secrets.token_urlsafe(32))
        except InputError:
            pass
        else:
            raise AssertionError("Wrong password accepted")
        assert sha(restored) == before
        return {
            "backup_sha256": hashlib.sha256(raw).hexdigest(),
            "restart_and_restore": True,
        }

    record(
        "Shared real encrypted persistence backup restore wrong-password preservation",
        persist,
    )


def main():
    global APP
    APP = Application(
        ROOT,
        BASE / "persistent-data",
        passphrase=PASSWORD,
    )
    try:
        recovery()
        shared()
    finally:
        APP.close()
        hashes = {
            p.relative_to(BASE).as_posix(): sha(p)
            for p in BASE.rglob("*")
            if p.is_file() and p.name != "hashes.json"
        }
        (BASE / "hashes.json").write_text(
            json.dumps(hashes, indent=2), encoding="utf-8"
        )
    return int(any(r["status"] == "FAIL" for r in RESULTS))


if __name__ == "__main__":
    raise SystemExit(main())
